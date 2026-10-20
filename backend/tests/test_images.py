import base64
import io

from PIL import Image
import pptx
import pytest
from pydantic import ValidationError

from backend.app.models.presentation import ElementModel, PresentationModel, SlideModel
from backend.app.services.image_assets import rendered_image
from backend.app.services.pptx_exporter import PPTXExporter
from backend.app.services.video_exporter import VideoExporter


def image_element(**patch):
    image = Image.new('RGB', (100, 50), 'red')
    image.paste('blue', (50, 0, 100, 50))
    stream = io.BytesIO()
    image.save(stream, format='PNG')
    return ElementModel.model_validate(dict(id='picture', type='image', x=20, y=20, width=100, height=50,
        image_width=100, image_height=50, image_src='data:image/png;base64,' + base64.b64encode(stream.getvalue()).decode(), **patch))


def test_image_crop_flip_opacity_and_color_adjustments():
    element = image_element(image_crop=dict(x=.5, y=0, width=.5, height=1), opacity=.5)
    image = rendered_image(element, element.opacity)
    assert image.size == (50, 50)
    assert image.getpixel((10, 10)) == (0, 0, 255, 128)
    assert rendered_image(image_element(flip_x=True)).getpixel((0, 0))[:3] == (0, 0, 255)
    gray = rendered_image(image_element(image_saturation=0)).getpixel((0, 0))
    assert gray[0] == gray[1] == gray[2]
    contrast = rendered_image(image_element(image_contrast=0)).getpixel((0, 0))
    assert contrast[:3] == (128, 128, 128)
    assert rendered_image(image_element(image_brightness=0)).getpixel((0, 0))[:3] == (0, 0, 0)


def test_invalid_crop_and_remote_source_are_rejected():
    with pytest.raises(ValidationError):
        image_element(image_crop=dict(x=.8, width=.5))
    with pytest.raises(ValidationError):
        ElementModel(id='bad', type='image', image_src='https://example.com/image.png')


def test_images_survive_save_model_pptx_and_static_video():
    element = image_element(image_crop=dict(x=.5, width=.5))
    deck = PresentationModel(id='images', slides=[SlideModel(id='slide', elements=[element])])
    loaded = PresentationModel.model_validate_json(deck.model_dump_json())
    assert loaded.slides[0].elements[0].image_src == element.image_src
    video = VideoExporter().render_slide_image(loaded.slides[0])
    assert video.getpixel((50, 40))[:3] == (0, 0, 255)
    document = pptx.Presentation(PPTXExporter().export(loaded))
    picture = document.slides[0].shapes[-1]
    assert picture.shape_type == 13
    embedded = Image.open(io.BytesIO(picture.image.blob))
    assert embedded.getpixel((0, 0))[:3] == (0, 0, 255)


def test_portrait_pptx_uses_portrait_image_coordinates():
    deck = PresentationModel(id='portrait', aspect_ratio='9:16', slides=[SlideModel(id='slide', elements=[image_element()])])
    document = pptx.Presentation(PPTXExporter().export(deck))
    assert document.slide_width < document.slide_height
    picture = document.slides[0].shapes[-1]
    assert abs(picture.width / document.slide_width - 100 / 720) < .001
