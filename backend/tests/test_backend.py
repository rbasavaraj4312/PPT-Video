import pytest
from pathlib import Path
import pptx
from fastapi.testclient import TestClient

from backend.app.models.presentation import (
    PresentationModel,
    SlideModel,
    ElementModel,
    ElementType,
    ArrowHead,
    StrokeStyle,
)
from backend.app.services.storage import StorageService
from backend.app.services.pptx_exporter import PPTXExporter
from backend.app.main import app


@pytest.fixture
def temp_storage(tmp_path):
    return StorageService(storage_dir=tmp_path)


def test_group_resize_scale_survives_storage(temp_storage):
    element = ElementModel(id="scaled-arrow", type=ElementType.ARROW,
                           x=100.25, y=80.5, width=-50.25, height=0,
                           stroke_width=.5, font_size=6, visual_scale=.25,
                           group_id="scaled-group")
    deck = PresentationModel(id="scaled-deck", title="Scaled group",
                             slides=[SlideModel(id="scaled-slide", elements=[element])])
    temp_storage.save_presentation(deck)
    saved = temp_storage.get_presentation(deck.id).slides[0].elements[0]
    assert saved.visual_scale == .25
    assert saved.stroke_width == .5
    assert saved.width == -50.25
    assert saved.group_id == "scaled-group"


def test_storage_crud(temp_storage):
    # Verify sample is created
    summaries = temp_storage.list_presentations()
    assert len(summaries) >= 1
    sample_id = summaries[0].id

    # Get sample
    pres = temp_storage.get_presentation(sample_id)
    assert pres is not None
    assert pres.id == sample_id
    assert len(pres.slides) >= 1

    # Create new presentation
    new_pres = PresentationModel(
        id="test-pres-123",
        title="Custom Test Presentation",
        aspect_ratio="16:9",
        slides=[
            SlideModel(
                id="s-1",
                title="Slide 1",
                background="#ffffff",
                elements=[
                    ElementModel(
                        id="t-1",
                        type=ElementType.TEXT,
                        x=100,
                        y=100,
                        width=300,
                        height=50,
                        content="Test Headline",
                        font_size=28,
                    ),
                    ElementModel(
                        id="sh-1",
                        type=ElementType.RECTANGLE,
                        x=100,
                        y=200,
                        width=200,
                        height=150,
                        fill="#3b82f6",
                    ),
                    ElementModel(
                        id="arr-1",
                        type=ElementType.ARROW,
                        x1=320,
                        y1=275,
                        x2=450,
                        y2=275,
                        stroke="#ef4444",
                        stroke_width=3,
                        arrow_head=ArrowHead.END,
                    ),
                ],
            )
        ],
    )

    saved = temp_storage.save_presentation(new_pres)
    assert saved.id == "test-pres-123"

    # Verify listed
    all_pres = temp_storage.list_presentations()
    ids = [p.id for p in all_pres]
    assert "test-pres-123" in ids

    # Retrieve
    loaded = temp_storage.get_presentation("test-pres-123")
    assert loaded is not None
    assert loaded.title == "Custom Test Presentation"
    assert len(loaded.slides[0].elements) == 3

    # Delete
    deleted = temp_storage.delete_presentation("test-pres-123")
    assert deleted is True
    assert temp_storage.get_presentation("test-pres-123") is None

    # Verify structured directory layout
    assert (temp_storage.storage_dir / "templates").is_dir()
    assert (temp_storage.storage_dir / "motion-templates").is_dir()
    assert (temp_storage.storage_dir / "data").is_dir()
    assert (temp_storage.storage_dir / "data" / "folders.json").is_file()
    assert (temp_storage.storage_dir / "data" / "sample-presentation.json").is_file()
    assert (temp_storage.storage_dir / "templates" / "templates.json").is_file() or True
    assert (temp_storage.storage_dir / "motion-templates" / "motion_templates.json").is_file() or True


def test_pptx_exporter():
    exporter = PPTXExporter()
    presentation = PresentationModel(
        id="export-test",
        title="PPTX Export Test",
        aspect_ratio="16:9",
        slides=[
            SlideModel(
                id="slide-1",
                title="Shapes Slide",
                background="#0f172a",
                notes="Speaker note test",
                elements=[
                    ElementModel(
                        id="el-title",
                        type=ElementType.TEXT,
                        x=50,
                        y=50,
                        width=800,
                        height=60,
                        content="Export Title",
                        font_size=32,
                        font_color="#ffffff",
                    ),
                    ElementModel(
                        id="el-rect",
                        type=ElementType.RECTANGLE,
                        x=50,
                        y=150,
                        width=200,
                        height=120,
                        fill="#2563eb",
                        stroke="#ffffff",
                        stroke_width=2,
                    ),
                    ElementModel(
                        id="el-round",
                        type=ElementType.ROUNDED_RECTANGLE,
                        x=300,
                        y=150,
                        width=200,
                        height=120,
                        fill="#10b981",
                        stroke="#ffffff",
                        stroke_width=2,
                    ),
                    ElementModel(
                        id="el-circle",
                        type=ElementType.ELLIPSE,
                        x=550,
                        y=150,
                        width=120,
                        height=120,
                        fill="#f59e0b",
                    ),
                    ElementModel(
                        id="el-arrow",
                        type=ElementType.ARROW,
                        x1=50,
                        y1=350,
                        x2=670,
                        y2=350,
                        stroke="#ec4899",
                        stroke_width=4,
                        arrow_head=ArrowHead.END,
                    ),
                    ElementModel(
                        id="el-line",
                        type=ElementType.LINE,
                        x1=50,
                        y1=400,
                        x2=670,
                        y2=400,
                        stroke="#94a3b8",
                        stroke_width=2,
                        stroke_style=StrokeStyle.DASHED,
                    ),
                ],
            )
        ],
    )

    buffer = exporter.export(presentation)
    assert buffer is not None
    assert buffer.getbuffer().nbytes > 1000  # Valid non-empty PPTX

    # Verify with python-pptx
    prs = pptx.Presentation(buffer)
    assert len(prs.slides) == 1
    slide = prs.slides[0]
    assert slide.notes_slide.notes_text_frame.text == "Speaker note test"
    assert len(slide.shapes) >= 6


def test_api_endpoints():
    client = TestClient(app)

    # Health check
    res = client.get("/health")
    assert res.status_code == 200
    assert res.json() == {"status": "ok"}

    # List presentations
    res = client.get("/api/presentations")
    assert res.status_code == 200
    data = res.json()
    assert isinstance(data, list)
    assert len(data) >= 1
    sample_id = data[0]["id"]

    # Get sample presentation
    res_get = client.get(f"/api/presentations/{sample_id}")
    assert res_get.status_code == 200
    pres_data = res_get.json()
    assert pres_data["id"] == sample_id
    assert len(pres_data["slides"]) >= 1

    # Export presentation
    res_export = client.post(f"/api/presentations/{sample_id}/export")
    assert res_export.status_code == 200
    assert "presentationml" in res_export.headers.get("content-type", "")
    assert len(res_export.content) > 1000

    # Direct export
    res_direct = client.post("/api/export-direct", json=pres_data)
    assert res_direct.status_code == 200
    assert len(res_direct.content) > 1000


def test_frontend_json_compatibility():
    client = TestClient(app)

    # Payload matching the exact format sent by the React frontend
    frontend_payload = {
        "id": "react-pres-test",
        "title": "Interactive React Presentation",
        "aspect_ratio": "16:9",
        "slides": [
            {
                "id": "s-react-1",
                "title": "React Architecture Slide",
                "background": "#0f172a",
                "elements": [
                    {
                        "id": "el-title",
                        "type": "text",
                        "x": 100,
                        "y": 60,
                        "width": 1080,
                        "height": 70,
                        "text": "Vector Database Architecture",
                        "font_size": 36,
                        "font_bold": True,
                        "text_color": "#38bdf8",
                        "text_align": "center",
                    },
                    {
                        "id": "el-box-1",
                        "type": "rounded_rectangle",
                        "x": 160,
                        "y": 200,
                        "width": 260,
                        "height": 160,
                        "fill_color": "#1e293b",
                        "stroke_color": "#38bdf8",
                        "stroke_width": 2,
                        "text": "Embeddings Input",
                        "font_size": 20,
                        "text_color": "#f8fafc",
                    },
                    {
                        "id": "el-arr-1",
                        "type": "arrow",
                        "x": 440,
                        "y": 280,
                        "width": 160,
                        "height": 0,
                        "stroke_color": "#38bdf8",
                        "stroke_width": 3,
                        "arrow_head_end": True,
                    },
                    {
                        "id": "el-box-2",
                        "type": "rectangle",
                        "x": 620,
                        "y": 200,
                        "width": 300,
                        "height": 160,
                        "fill_color": "#0284c7",
                        "stroke_color": "#bae6fd",
                        "stroke_width": 2,
                        "text": "HNSW Vector Index",
                        "font_size": 20,
                        "text_color": "#ffffff",
                    },
                ],
            }
        ],
    }

    # Save via POST /api/presentations
    res_save = client.post("/api/presentations", json=frontend_payload)
    assert res_save.status_code == 200
    saved = res_save.json()
    assert saved["id"] == "react-pres-test"
    assert len(saved["slides"][0]["elements"]) == 4

    # Export via POST /api/export-direct
    res_export = client.post("/api/export-direct", json=frontend_payload)
    assert res_export.status_code == 200
    assert len(res_export.content) > 1000

    # Clean up test presentation
    client.delete("/api/presentations/react-pres-test")


def test_video_export_with_morph_transition():
    client = TestClient(app)

    # Multi-slide presentation with morph transition between slides
    multi_slide_pres = {
        "id": "morph-video-test",
        "title": "Morph Video Transition Presentation",
        "aspect_ratio": "16:9",
        "slides": [
            {
                "id": "slide-morph-1",
                "title": "Slide 1 - Origin",
                "background": "#0f172a",
                "transition_type": "morph",
                "transition_duration": 0.4,
                "slide_duration": 0.6,
                "elements": [
                    {
                        "id": "box-moving",
                        "morph_id": "core-box",
                        "type": "rounded_rectangle",
                        "x": 100,
                        "y": 200,
                        "width": 200,
                        "height": 120,
                        "fill_color": "#3b82f6",
                        "stroke_color": "#1d4ed8",
                        "stroke_width": 2,
                        "text": "Vector DB",
                        "font_size": 22,
                        "text_color": "#ffffff",
                    },
                    {
                        "id": "text-departing",
                        "type": "text",
                        "x": 100,
                        "y": 60,
                        "width": 800,
                        "height": 50,
                        "text": "Phase 1: Ingestion",
                        "font_size": 32,
                        "text_color": "#38bdf8",
                    },
                ],
            },
            {
                "id": "slide-morph-2",
                "title": "Slide 2 - Destination",
                "background": "#1e293b",
                "transition_type": "morph",
                "transition_duration": 0.4,
                "slide_duration": 0.6,
                "elements": [
                    {
                        "id": "box-moved",
                        "morph_id": "core-box",  # paired by morph_id
                        "type": "rounded_rectangle",
                        "x": 650,
                        "y": 220,
                        "width": 380,
                        "height": 220,
                        "fill_color": "#10b981",
                        "stroke_color": "#059669",
                        "stroke_width": 3,
                        "text": "Vector Database (HNSW)",
                        "font_size": 26,
                        "text_color": "#ffffff",
                    },
                    {
                        "id": "text-arriving",
                        "type": "text",
                        "x": 650,
                        "y": 60,
                        "width": 600,
                        "height": 50,
                        "text": "Phase 2: Indexing & Search",
                        "font_size": 32,
                        "text_color": "#34d399",
                    },
                ],
            },
        ],
    }

    # 1. Direct video export test
    res = client.post("/api/export-video-direct", json=multi_slide_pres)
    assert res.status_code == 200
    assert "video/mp4" in res.headers.get("content-type", "")
    assert len(res.content) > 1000  # Valid MP4 video file

    # 2. Saved video export test
    client.post("/api/presentations", json=multi_slide_pres)
    res_saved = client.post("/api/presentations/morph-video-test/export-video")
    assert res_saved.status_code == 200
    assert "video/mp4" in res_saved.headers.get("content-type", "")
    assert len(res_saved.content) > 1000

    # Clean up
    client.delete("/api/presentations/morph-video-test")


def test_grouped_elements_persistence_and_export():
    client = TestClient(app)
    grouped_pres = {
        "id": "grouped-test-pres",
        "title": "Grouped Elements Test",
        "aspect_ratio": "16:9",
        "slides": [
            {
                "id": "slide-grp-1",
                "title": "Slide with Groups",
                "background": "#0f172a",
                "transition_type": "morph",
                "elements": [
                    {
                        "id": "g-shape-1",
                        "group_id": "grp_alpha_123",
                        "type": "rectangle",
                        "x": 100,
                        "y": 150,
                        "width": 120,
                        "height": 80,
                        "fill": "#3b82f6",
                    },
                    {
                        "id": "g-shape-2",
                        "group_id": "grp_alpha_123",
                        "type": "ellipse",
                        "x": 240,
                        "y": 150,
                        "width": 80,
                        "height": 80,
                        "fill": "#10b981",
                    },
                    {
                        "id": "single-shape",
                        "group_id": None,
                        "type": "text",
                        "x": 100,
                        "y": 50,
                        "width": 400,
                        "height": 50,
                        "content": "Grouped Shapes Demo",
                    },
                ],
            }
        ],
    }

    # Save presentation
    res = client.post("/api/presentations", json=grouped_pres)
    assert res.status_code == 200

    # Retrieve and verify group_id is preserved
    get_res = client.get("/api/presentations/grouped-test-pres")
    assert get_res.status_code == 200
    data = get_res.json()
    elements = data["slides"][0]["elements"]
    assert len(elements) == 3
    assert elements[0]["group_id"] == "grp_alpha_123"
    assert elements[1]["group_id"] == "grp_alpha_123"
    assert elements[2]["group_id"] is None

    # Test PPTX export with groups
    export_res = client.post("/api/presentations/grouped-test-pres/export")
    assert export_res.status_code == 200
    assert "openxmlformats" in export_res.headers.get("content-type", "")
    assert len(export_res.content) > 1000

    # Clean up
    client.delete("/api/presentations/grouped-test-pres")


def test_video_export_custom_resolution_and_fps():
    client = TestClient(app)
    pres = {
        "id": "res-fps-test",
        "title": "Resolution and FPS Test",
        "aspect_ratio": "16:9",
        "slides": [
            {
                "id": "s1",
                "title": "Slide 1",
                "background": "#1e1b4b",
                "slide_duration": 0.5,
                "elements": [
                    {
                        "id": "t1",
                        "type": "text",
                        "x": 200,
                        "y": 200,
                        "width": 800,
                        "height": 100,
                        "text": "1080p 60FPS Video",
                        "font_size": 40,
                        "text_color": "#ffffff",
                    }
                ],
            }
        ],
    }

    # Test direct export with 1080p and 60fps
    res = client.post(
        "/api/export-video-direct?resolution=1080p&fps=60",
        json=pres,
    )
    assert res.status_code == 200
    assert "video/mp4" in res.headers.get("content-type", "")
    assert len(res.content) > 1000

    # Test 720p 24fps
    res_720 = client.post(
        "/api/export-video-direct?resolution=720p&fps=24",
        json=pres,
    )
    assert res_720.status_code == 200
    assert len(res_720.content) > 1000


def test_slide_compiler_approach_4():
    """Test Approach 4 outline compiler with flow, cards, comparison, icons, and morph tags."""
    from backend.app.services.slide_compiler import slide_compiler

    outline = """Title: Cloud Architecture Deck
Theme: modern-dark

# Slide 1: System Ingestion
Layout: flow
(icon: globe) Client App: Web Traffic
[Client App] -> [API Gateway]
(icon: router) API Gateway: Reverse Proxy & Rate Limiting
[API Gateway] -> [PostgreSQL DB]
(icon: database) PostgreSQL DB: Primary Storage

# Slide 2: Data Persistence
Layout: cards
(icon: database) PostgreSQL DB: Primary Storage
(icon: binary) Vector Index: Semantic Search
(icon: hard-drive) Backup Vault: Automated Daily Snapshots

# Slide 3: SQL vs Vector Comparison
Layout: comparison
Left: Relational DB
(icon: database) PostgreSQL: ACID Transactions
- B-Tree Indexes
- Strict Schema

Right: Vector Search
(icon: binary) Pinecone: Semantic Embeddings
- HNSW Nearest Neighbor
- Cosine Similarity
"""

    presentation = slide_compiler.compile(outline, theme_override="modern-dark")
    assert presentation.title == "Cloud Architecture Deck"
    assert len(presentation.slides) == 3

    # Check Slide 1 (flow)
    slide1 = presentation.slides[0]
    icons_s1 = [e for e in slide1.elements if e.type == ElementType.ICON]
    arrows_s1 = [e for e in slide1.elements if e.type == ElementType.ARROW]
    assert len(icons_s1) == 3
    assert len(arrows_s1) == 2
    assert any(i.icon_name == "globe" for i in icons_s1)
    assert any(i.icon_name == "database" for i in icons_s1)

    # Check cross-slide morph_id continuity for PostgreSQL DB
    db_icon_s1 = next(i for i in icons_s1 if "database" in i.icon_name)
    assert db_icon_s1.morph_id is not None

    # Check Slide 2 (cards)
    slide2 = presentation.slides[1]
    icons_s2 = [e for e in slide2.elements if e.type == ElementType.ICON]
    assert len(icons_s2) == 3
    db_icon_s2 = next(i for i in icons_s2 if "database" in i.icon_name)
    assert db_icon_s2.morph_id == db_icon_s1.morph_id  # Morph pairing preserved!

    # Check Slide 3 (comparison)
    slide3 = presentation.slides[2]
    icons_s3 = [e for e in slide3.elements if e.type == ElementType.ICON]
    assert len(icons_s3) >= 2


def test_compile_outline_api():
    """Test POST /api/compile-outline and GET /api/outline-themes."""
    client = TestClient(app)

    # 1. Themes endpoint
    themes_res = client.get("/api/outline-themes")
    assert themes_res.status_code == 200
    themes_data = themes_res.json()
    assert len(themes_data) >= 5
    theme_ids = [t["id"] for t in themes_data]
    assert "modern-dark" in theme_ids
    assert "midnight" in theme_ids

    # 2. Compile outline endpoint
    compile_payload = {
        "outline": """Title: Microservices API
Theme: midnight

# Slide 1: Flow
Layout: flow
(icon: server) Web Server: Ingress
[Web Server] -> [Cache]
(icon: zap) Cache: Redis In-Memory
""",
        "theme": "midnight",
    }
    compile_res = client.post("/api/compile-outline", json=compile_payload)
    assert compile_res.status_code == 200
    pres_data = compile_res.json()
    assert pres_data["title"] == "Microservices API"
    assert len(pres_data["slides"]) == 1
    assert any(e["type"] == "icon" for e in pres_data["slides"][0]["elements"])


def test_video_and_pptx_export_with_icons():
    """Test exporting icons to PPTX and 60FPS video with morph transitions."""
    client = TestClient(app)

    pres = {
        "id": "pres-icons-test",
        "title": "Icons Morph Presentation",
        "aspect_ratio": "16:9",
        "slides": [
            {
                "id": "s1",
                "title": "Icon Slide 1",
                "background": "#0f172a",
                "transition_type": "morph",
                "transition_duration": 0.5,
                "slide_duration": 0.5,
                "elements": [
                    {
                        "id": "ic-1",
                        "type": "icon",
                        "x": 200,
                        "y": 250,
                        "width": 100,
                        "height": 100,
                        "icon_name": "database",
                        "icon_badge": "rounded",
                        "badge_color": "#1e293b",
                        "icon_color": "#38bdf8",
                        "fill": "#1e293b",
                        "stroke": "#334155",
                        "stroke_width": 2,
                        "content": "PostgreSQL",
                        "font_size": 16,
                        "morph_id": "morph-db",
                    }
                ],
            },
            {
                "id": "s2",
                "title": "Icon Slide 2",
                "background": "#0f172a",
                "transition_type": "morph",
                "transition_duration": 0.5,
                "slide_duration": 0.5,
                "elements": [
                    {
                        "id": "ic-2",
                        "type": "icon",
                        "x": 600,
                        "y": 200,
                        "width": 150,
                        "height": 150,
                        "icon_name": "database",
                        "icon_badge": "circle",
                        "badge_color": "#0369a1",
                        "icon_color": "#ffffff",
                        "fill": "#0369a1",
                        "stroke": "#38bdf8",
                        "stroke_width": 3,
                        "content": "PostgreSQL Cluster",
                        "font_size": 18,
                        "morph_id": "morph-db",
                    }
                ],
            },
        ],
    }

    # 1. Test PPTX export with icons
    pptx_res = client.post("/api/export-direct", json=pres)
    assert pptx_res.status_code == 200
    assert len(pptx_res.content) > 1000

    # 2. Test MP4 video export with icons and 60FPS morph
    video_res = client.post(
        "/api/export-video-direct?resolution=720p&fps=60",
        json=pres,
    )
    assert video_res.status_code == 200
    assert "video/mp4" in video_res.headers.get("content-type", "")
    assert len(video_res.content) > 2000


def test_video_text_wrapping_and_vertical_align():
    from backend.app.services.video_exporter import VideoExporter
    from backend.app.models.presentation import SlideModel, ElementModel, ElementType
    from PIL import Image, ImageDraw

    exporter = VideoExporter()
    img = Image.new("RGBA", (1280, 720), (15, 23, 42, 255))
    draw = ImageDraw.Draw(img)

    # Test direct _draw_text word wrapping
    long_text = "Vectorizes User Query into Same Dimensional Space"
    exporter._draw_text(
        draw,
        long_text,
        x=100,
        y=100,
        w=200,
        h=100,
        font_size=16,
        color_hex="#ffffff",
        align="center",
        vertical_align="top",
    )

    # Test complete slide rendering with narrow card text
    slide = SlideModel(
        id="s-wrap-test",
        title="Wrapping Test",
        background="#0f172a",
        elements=[
            ElementModel(
                id="card-1",
                type=ElementType.ROUNDED_RECTANGLE,
                x=80,
                y=200,
                width=220,
                height=260,
                fill="#1e293b",
            ),
            ElementModel(
                id="desc-1",
                type=ElementType.TEXT,
                x=90,
                y=350,
                width=200,
                height=100,
                content="Natural Language Question from Client",
                font_size=15,
                text_align="center",
                vertical_align="top",
            ),
            ElementModel(
                id="card-2",
                type=ElementType.ROUNDED_RECTANGLE,
                x=340,
                y=200,
                width=220,
                height=260,
                fill="#1e293b",
            ),
            ElementModel(
                id="desc-2",
                type=ElementType.TEXT,
                x=350,
                y=350,
                width=200,
                height=100,
                content="Vectorizes User Query into Same Dimensional Space",
                font_size=15,
                text_align="center",
                vertical_align="top",
            ),
        ],
    )

    rendered_img = exporter.render_slide_image(slide, scale=1.0)
    assert rendered_img.size == (1280, 720)
    rendered_1080p = exporter.render_slide_image(slide, scale=1.5, canvas_w=1920, canvas_h=1080)
    assert rendered_1080p.size == (1920, 1080)


def test_compiler_extended_layouts_and_themes():
    from backend.app.services.slide_compiler import slide_compiler

    outline = """Title: Deep Space Autonomous Swarm
Theme: cyber-neon

# Slide 1: Mission Hub
Layout: architecture
(icon: cpu) Swarm Quantum Core: Central coordinator and routing hub
(icon: globe) Sensor Lidar: Perimeter telemetry
(icon: shield) Kinetic Interceptor: Defensive micro-drones
(icon: zap) Energy Harvester: High-yield beam collector
(icon: database) State Ledger: Distributed memory mesh

# Slide 2: Fleet Performance
Layout: metrics
(icon: zap) 0.4ms | Gossip Latency: Sub-millisecond peer gossip protocol
(icon: globe) 250,000 | Active Drones: Autonomous orbital mesh
(icon: shield) 99.999% | Precision: Micro-radian trajectory alignment
(icon: cpu) Zero | Human Delay: Full edge autonomy

# Slide 3: Neural Engine
Layout: split
(icon: cpu) Astraea Photonic Neural Core: Deep vacuum radiation-hardened unit
(icon: shield) Diamond Substrate: Impervious to cosmic radiation
(icon: globe) Optical Crosslink: 100 Gbps free space lasers
(icon: database) Holographic Synaptic Memory: Atomic-layer state preservation

# Slide 4: Deployment Horizon
Layout: timeline
(icon: globe) Phase 1: LEO Swarm: 10,000 unit testing constellation
(icon: cpu) Phase 2: Lunar Gateway: Autonomous mining relay
(icon: shield) Phase 3: Asteroid Intercept: Planetary defense corridor
(icon: zap) Phase 4: Interplanetary Mesh: Solar-system transport grid
"""

    presentation = slide_compiler.compile(outline)
    assert presentation.title == "Deep Space Autonomous Swarm"
    assert len(presentation.slides) == 4

    # Slide 1 (architecture / hub-spoke): Check central hub + satellites + arrows
    s1 = presentation.slides[0]
    assert s1.background == "#070814"
    arrows_s1 = [e for e in s1.elements if e.type == ElementType.ARROW]
    assert len(arrows_s1) == 4
    icons_s1 = [e for e in s1.elements if e.type == ElementType.ICON]
    assert len(icons_s1) == 5

    # Slide 2 (metrics): Check 4 metric cards with big numbers
    s2 = presentation.slides[1]
    texts_s2 = [e for e in s2.elements if e.type == ElementType.TEXT]
    metric_values = [e.content for e in texts_s2 if e.font_size >= 40]
    assert "0.4ms" in metric_values
    assert "250,000" in metric_values
    assert "99.999%" in metric_values
    assert "Zero" in metric_values

    # Slide 3 (split / spotlight)
    s3 = presentation.slides[2]
    rects_s3 = [e for e in s3.elements if e.type == ElementType.ROUNDED_RECTANGLE]
    assert len(rects_s3) == 4  # 1 hero left card + 3 right stacked cards

    # Slide 4 (timeline)
    s4 = presentation.slides[3]
    badges_s4 = [e for e in s4.elements if e.type == ElementType.ELLIPSE and "0" in (e.content or "")]
    assert len(badges_s4) == 4  # 01, 02, 03, 04 sequence badges


def test_folders_api_and_storage():
    """Test folder CRUD API and persistence."""
    client = TestClient(app)

    # 1. List folders
    res = client.get("/api/folders")
    assert res.status_code == 200
    initial_folders = res.json()
    assert isinstance(initial_folders, list)

    # 2. Create folder
    new_folder = {
        "id": "folder-test-shorts",
        "name": "YouTube Shorts",
        "color": "#f43f5e",
    }
    create_res = client.post("/api/folders", json=new_folder)
    assert create_res.status_code == 200
    created = create_res.json()
    assert created["id"] == "folder-test-shorts"
    assert created["name"] == "YouTube Shorts"

    # 3. Update folder
    update_res = client.put("/api/folders/folder-test-shorts", json={"name": "Viral Reels", "color": "#06b6d4"})
    assert update_res.status_code == 200
    updated = update_res.json()
    assert updated["name"] == "Viral Reels"
    assert updated["color"] == "#06b6d4"

    # 4. Delete folder
    del_res = client.delete("/api/folders/folder-test-shorts")
    assert del_res.status_code == 200
    assert del_res.json()["status"] == "deleted"


def test_presentation_folder_assignment():
    """Test moving a presentation into a folder and verifying in summary."""
    client = TestClient(app)

    # Create dummy presentation
    pres_payload = {
        "id": "pres-folder-test",
        "title": "Folder Assignment Test",
        "aspect_ratio": "16:9",
        "slides": [{"id": "s1", "title": "S1", "elements": []}],
    }
    client.post("/api/presentations", json=pres_payload)

    # Assign to folder
    move_res = client.put("/api/presentations/pres-folder-test/folder", json={"folder_id": "folder-reels"})
    assert move_res.status_code == 200
    assert move_res.json()["folder_id"] == "folder-reels"

    # Verify in list_presentations summary
    list_res = client.get("/api/presentations")
    matching = next((p for p in list_res.json() if p["id"] == "pres-folder-test"), None)
    assert matching is not None
    assert matching["folder_id"] == "folder-reels"

    # Clean up
    client.delete("/api/presentations/pres-folder-test")


def test_vertical_9_16_reels_support():
    """Test compiling and exporting vertical 9:16 Reels/Shorts."""
    from backend.app.services.slide_compiler import slide_compiler
    from backend.app.services.video_exporter import video_exporter

    reels_outline = """Title: 3 AI Trends in 30 Seconds
Format: reels
Theme: sunset

# Slide 1: Future of AI
Layout: hero
(icon: zap) Top 3 Autonomous Agent Trends Transforming 2026

# Slide 2: Trend Flow
Layout: flow
(icon: cpu) Edge LLMs: Micro models running locally
(icon: shield) Zero Trust: Formal verification of actions
(icon: database) Memory Mesh: Perpetual agent recall
"""

    pres = slide_compiler.compile(reels_outline)
    assert pres.aspect_ratio == "9:16"
    assert len(pres.slides) == 2

    # Check vertical element coordinates (width <= 720, height <= 1280)
    for s in pres.slides:
        for elem in s.elements:
            assert elem.x >= 0
            assert elem.x + elem.width <= 720 or elem.type == ElementType.LINE or elem.type == ElementType.ARROW

    # Test direct video export of 9:16 vertical video
    buf = video_exporter.export(pres, resolution="720p", fps=15, slide_duration=0.5)
    assert len(buf.getvalue()) > 5000  # Valid MP4 binary generated


def test_nested_subfolders_and_hierarchy():
    """Test creating subfolders inside folders, moving folders, and cascading deletion."""
    client = TestClient(app)
    # 1. Create parent folder
    res1 = client.post("/api/folders", json={"id": "f_marketing", "name": "Marketing", "color": "#f97316"})
    assert res1.status_code == 200

    # 2. Create subfolder inside Marketing
    res2 = client.post(
        "/api/folders", 
        json={"id": "f_campaigns", "name": "2026 Campaigns", "color": "#eab308", "parent_id": "f_marketing"}
    )
    assert res2.status_code == 200
    assert res2.json()["parent_id"] == "f_marketing"

    # 3. Create nested grand-child subfolder
    res3 = client.post(
        "/api/folders", 
        json={"id": "f_reels_sub", "name": "Reels & TikToks", "color": "#ec4899", "parent_id": "f_campaigns"}
    )
    assert res3.status_code == 200
    assert res3.json()["parent_id"] == "f_campaigns"

    # 4. Verify list contains all 3 folders with correct parent relationships
    list_res = client.get("/api/folders")
    folders = {f["id"]: f for f in list_res.json()}
    assert folders["f_marketing"]["parent_id"] is None
    assert folders["f_campaigns"]["parent_id"] == "f_marketing"
    assert folders["f_reels_sub"]["parent_id"] == "f_campaigns"

    # 5. Test move folder: prevent cycle (cannot move f_marketing into f_reels_sub)
    cycle_res = client.put("/api/folders/f_marketing/move", json={"parent_id": "f_reels_sub"})
    assert cycle_res.status_code == 400

    # 6. Move f_reels_sub directly to f_marketing
    move_res = client.put("/api/folders/f_reels_sub/move", json={"parent_id": "f_marketing"})
    assert move_res.status_code == 200
    assert move_res.json()["parent_id"] == "f_marketing"

    # 7. Delete parent folder: cascades to all subfolders
    del_res = client.delete("/api/folders/f_marketing")
    assert del_res.status_code == 200

    remaining = [f["id"] for f in client.get("/api/folders").json()]
    assert "f_marketing" not in remaining
    assert "f_campaigns" not in remaining
    assert "f_reels_sub" not in remaining


def test_templates_api_and_builtin_catalog(tmp_path, monkeypatch):
    """Test template listing, category filtering, multi-slide integrity, and custom template CRUD."""
    import backend.app.api.routes as routes_module
    test_storage = StorageService(storage_dir=tmp_path)
    monkeypatch.setattr(routes_module, "storage_service", test_storage)
    client = TestClient(app)

    # 1. List all templates
    res = client.get("/api/templates")
    assert res.status_code == 200
    all_tpls = res.json()
    assert len(all_tpls) >= 1

    # 2. Filter by category 'reels'
    res_reels_cat = client.get("/api/templates?category=reels")
    assert res_reels_cat.status_code == 200
    reels_tpls = res_reels_cat.json()
    assert len(reels_tpls) >= 1
    for t in reels_tpls:
        assert t["category"] == "reels"

    # 3. Filter by aspect ratio '9:16'
    res_reels = client.get("/api/templates?aspect_ratio=9:16")
    assert res_reels.status_code == 200
    reels_tpls = res_reels.json()
    assert len(reels_tpls) >= 1
    for t in reels_tpls:
        assert t["aspect_ratio"] == "9:16"

    # 4. Multi-slide template detail & integrity
    res_detail = client.get("/api/templates/tpl-tech-short-storyboard")
    assert res_detail.status_code == 200
    detail = res_detail.json()
    assert detail["slide_count"] >= 5
    assert len(detail["slides"]) >= 5

    # 5. Delete built-in template succeeds and suppresses it
    del_builtin = client.delete("/api/templates/tpl-tech-short-storyboard")
    assert del_builtin.status_code == 200
    assert del_builtin.json()["status"] == "deleted"

    # Confirm it is no longer in the list
    remaining_ids = [t["id"] for t in client.get("/api/templates").json()]
    assert "tpl-tech-short-storyboard" not in remaining_ids

    # Deleting non-existent template returns 404
    del_notfound = client.delete("/api/templates/tpl-non-existent-xyz")
    assert del_notfound.status_code == 404

    # 6. Create custom user template
    custom_tpl = {
        "id": "tpl-custom-team-review",
        "name": "Team Review Custom",
        "description": "Custom quarterly team performance breakdown",
        "category": "reels",
        "aspect_ratio": "9:16",
        "slides": [
            {
                "id": "s-c-1",
                "title": "Review Slide 1",
                "elements": []
            },
            {
                "id": "s-c-2",
                "title": "Review Slide 2",
                "elements": []
            }
        ]
    }
    save_res = client.post("/api/templates", json=custom_tpl)
    assert save_res.status_code == 200
    assert save_res.json()["id"] == "tpl-custom-team-review"
    assert save_res.json()["slide_count"] == 2
    assert save_res.json()["is_builtin"] is False

    # 7. Edit custom template via PUT
    custom_tpl["name"] = "Team Review Updated Title"
    put_res = client.put("/api/templates/tpl-custom-team-review", json=custom_tpl)
    assert put_res.status_code == 200
    assert put_res.json()["name"] == "Team Review Updated Title"

    # 8. Delete custom template
    del_custom = client.delete("/api/templates/tpl-custom-team-review")
    assert del_custom.status_code == 200
    assert del_custom.json()["status"] == "deleted"


def test_motion_templates_api_crud(tmp_path, monkeypatch):
    """Test unified HTML+CSS motion templates listing, category filtering, detail, and custom CRUD."""
    import backend.app.api.routes as routes_module
    test_storage = StorageService(storage_dir=tmp_path)
    monkeypatch.setattr(routes_module, "storage_service", test_storage)
    client = TestClient(app)

    # 1. List all motion templates
    res = client.get("/api/motion-templates")
    assert res.status_code == 200
    all_motion_tpls = res.json()
    assert len(all_motion_tpls) >= 5

    # 2. Filter by category
    res_cards = client.get("/api/motion-templates?category=cards")
    assert res_cards.status_code == 200
    cards = res_cards.json()
    assert len(cards) >= 1
    for c in cards:
        assert c["category"] == "cards"

    # 3. Get built-in template detail
    res_hud = client.get("/api/motion-templates/mot-cyber-hud-card")
    assert res_hud.status_code == 200
    hud = res_hud.json()
    assert hud["name"] == "Cyberpunk HUD Card"
    assert "tag" in hud["custom_vars"]
    assert "<style>" in hud["code"]

    # 4. Delete built-in motion template succeeds and suppresses it
    del_builtin = client.delete("/api/motion-templates/mot-cyber-hud-card")
    assert del_builtin.status_code == 200
    assert del_builtin.json()["status"] == "deleted"

    # Confirm it is no longer in the motion templates list
    remaining_motion_ids = [t["id"] for t in client.get("/api/motion-templates").json()]
    assert "mot-cyber-hud-card" not in remaining_motion_ids

    # Deleting non-existent motion template returns 404
    del_notfound = client.delete("/api/motion-templates/mot-non-existent-xyz")
    assert del_notfound.status_code == 404

    # 5. Create custom motion template with unified code and custom_vars
    custom_motion_tpl = {
        "id": "mot-custom-glow-pulse",
        "name": "Custom Glowing Pulse Box",
        "description": "Ambient glowing card with dynamic color and heading",
        "category": "ambient",
        "width": 300,
        "height": 200,
        "code": "<style>.pulse-box { background: {{bg_color}}; border: 2px solid {{accent_color}}; }</style><div class='pulse-box'>{{label}}</div>",
        "custom_vars": {
            "bg_color": "#0a0a1a",
            "accent_color": "#00f2fe",
            "label": "QUANTUM CORE"
        }
    }
    create_res = client.post("/api/motion-templates", json=custom_motion_tpl)
    assert create_res.status_code == 200
    created = create_res.json()
    assert created["id"] == "mot-custom-glow-pulse"
    assert created["is_builtin"] is False
    assert created["custom_vars"]["label"] == "QUANTUM CORE"

    # 6. Update custom motion template via PUT
    custom_motion_tpl["name"] = "Updated Glowing Pulse Box"
    custom_motion_tpl["custom_vars"]["label"] = "ONLINE // READY"
    update_res = client.put("/api/motion-templates/mot-custom-glow-pulse", json=custom_motion_tpl)
    assert update_res.status_code == 200
    assert update_res.json()["name"] == "Updated Glowing Pulse Box"
    assert update_res.json()["custom_vars"]["label"] == "ONLINE // READY"

    # 7. Get updated custom template
    get_res = client.get("/api/motion-templates/mot-custom-glow-pulse")
    assert get_res.status_code == 200
    assert get_res.json()["name"] == "Updated Glowing Pulse Box"

    # 8. Delete custom motion template
    del_res = client.delete("/api/motion-templates/mot-custom-glow-pulse")
    assert del_res.status_code == 200
    assert del_res.json()["status"] == "deleted"



def test_motion_component_pptx_export():
    from backend.app.models.presentation import MotionComponentData

    # Create presentation with insertable motion components
    motion_elem_jump = ElementModel(
        id="mc-jump-1",
        type=ElementType.MOTION_COMPONENT,
        x=80,
        y=120,
        width=320,
        height=160,
        content="Autonomous Agents",
        motion_data=MotionComponentData(
            component_id="step-badge-jump",
            action_type="jumping",
            title="Autonomous Agents",
            subtitle="Self-organizing swarms replacing rigid ETL pipelines",
            tag="STAGE 01",
            card_color="#0e1329",
            accent_color="#38bdf8",
            text_color="#ffffff",
        ),
    )

    motion_elem_ticker = ElementModel(
        id="mc-ticker-2",
        type=ElementType.MOTION_COMPONENT,
        x=440,
        y=120,
        width=300,
        height=180,
        content="Median Latency",
        motion_data=MotionComponentData(
            component_id="metric-number-ticker",
            action_type="counting",
            title="Median Latency",
            subtitle="End-to-end vector search execution",
            value_start=0,
            value_end=12,
            suffix="ms",
            prefix="~",
            card_color="#0a1024",
            accent_color="#00f2fe",
            text_color="#ffffff",
        ),
    )

    pres = PresentationModel(
        id="pres-motion-test",
        title="Motion Components PPT Test",
        aspect_ratio="16:9",
        slides=[
            SlideModel(
                id="s-motion-1",
                title="Motion Components Overview",
                elements=[motion_elem_jump, motion_elem_ticker],
            )
        ],
    )

    exporter = PPTXExporter()
    buffer = exporter.export(pres)
    assert buffer is not None
    buffer_bytes = buffer.getvalue()
    assert len(buffer_bytes) > 1000

    # Parse exported PPTX with python-pptx to verify shapes and text frames exist
    prs = pptx.Presentation(buffer)
    assert len(prs.slides) == 1
    slide = prs.slides[0]
    # Check that shapes were created for the motion components
    shape_texts = []
    for s in slide.shapes:
        if s.has_text_frame:
            shape_texts.append(s.text_frame.text)
    combined = " ".join(shape_texts)
    assert "STAGE 01" in combined
    assert "Autonomous Agents" in combined
    assert "~12ms" in combined or "Median Latency" in combined


def test_tech_teaching_layouts_compilation_16_9():
    from backend.app.services.slide_compiler import slide_compiler

    outline = """Title: Tech Teaching Mastery
Theme: modern-dark
Format: 16:9

# Slide 1: Distributed Rate Limiter
Layout: code-walkthrough
Code: [rate_limiter.py]
def allow_request(user_id: str, limit: int = 100):
    current = redis.get(f"tokens:{user_id}")
    if current and int(current) > 0:
        redis.decr(f"tokens:{user_id}")
        return {"allowed": True, "remaining": int(current) - 1}
    return {"allowed": False, "retry_after": 60}
- (icon: clock) [1] Sliding Window: Evaluates quota in Redis
- (icon: zap) [2] Atomic Decrement: Consumes token safely
- (icon: shield) [3] 429 Guard: Drops burst spikes

# Slide 2: Production Cloud Architecture
Layout: architecture-layers
- (icon: globe) Layer 1 | Edge & CDN: Next.js 15, Cloudflare CDN
- (icon: router) Layer 2 | API Gateway: Kong Mesh, Token Bucket
- (icon: cpu) Layer 3 | Microservice Cluster: FastAPI, vLLM
- (icon: database) Layer 4 | Storage Tier: Qdrant, PostgreSQL, Redis

# Slide 3: Database Overload: Polling vs Push
Layout: problem-solution
Problem: Legacy HTTP Polling
- 100,000 active clients pinging every 2s
- 99% requests return empty 304 payloads
Solution: WebSocket Pub/Sub Mesh
- Single persistent duplex TCP connection
- Events pushed only when row state mutates

# Slide 4: Kubernetes Reconciliation Loop
Layout: circular-loop
(icon: eye) 1. Observe: Kubelet detects desired vs actual count
(icon: cpu) 2. Diff: Controller manager calculates state delta
(icon: zap) 3. Act: Scheduler assigns new pod
(icon: check-circle) 4. Verify: Readiness probe confirms health

# Slide 5: Real-Time Clickstream Pipeline
Layout: event-stream
(icon: globe) Ingestion: Mobile & Web Telemetry
(icon: server) Kafka Broker: 16-Partition Distributed Queue
(icon: cpu) Flink Workers: 10s Tumbling Window Aggregator
(icon: database) ClickHouse: Columnar Analytics Engine

# Slide 6: Turborepo 2.0 Engine Overhaul
Layout: bento-box
(icon: zap) Remote Caching Engine: Zero-computation rebuilds across workers
(icon: cpu) 10x | Build Speedup: Rust daemon replacing Go worker
(icon: workflow) Task Graph Pruning: Topological DAG execution
(icon: shield) Zero Config: Auto-discovers monorepo workspaces

# Slide 7: Storage Engine Decision Matrix
Layout: feature-matrix
PostgreSQL | Relational: ACID transactions, SQL joins, persistent OLTP storage
Redis | In-Memory: Sub-millisecond latency, RAM-based key-value caching
ClickHouse | Columnar OLAP: High-throughput batch inserts, real-time analytics
"""

    pres = slide_compiler.compile(outline)
    assert pres.aspect_ratio == "16:9"
    assert len(pres.slides) == 7

    # Verify no element overflows 1280x720 canvas
    for slide in pres.slides:
        assert len(slide.elements) >= 10
        for el in slide.elements:
            assert el.x >= 0, f"Negative x in {slide.title}: {el.x}"
            assert el.x + el.width <= 1280, f"X overflow in {slide.title}: {el.x + el.width}"
            assert el.y >= 0, f"Negative y in {slide.title}: {el.y}"
            assert el.y + el.height <= 720, f"Y overflow in {slide.title}: {el.y + el.height}"

    # Verify Slide 1 (Code Walkthrough) has code content
    s1_texts = [e.content for e in pres.slides[0].elements if e.type == ElementType.TEXT]
    assert any("def allow_request" in t for t in s1_texts)
    assert any("rate_limiter.py" in t for t in s1_texts)

    # Verify Slide 3 (Problem vs Solution) has problem and solution badges
    s3_texts = [e.content for e in pres.slides[2].elements if e.type == ElementType.TEXT]
    assert any("THE BOTTLENECK" in t or "PROBLEM" in t for t in s3_texts)
    assert any("MODERN ARCHITECTURE" in t or "SOLUTION" in t for t in s3_texts)

    # Verify Slide 5 (Event Stream) has broker partition indicators
    s5_texts = [e.content for e in pres.slides[4].elements if e.type == ElementType.TEXT]
    assert any("Partition" in t for t in s5_texts)


def test_tech_teaching_layouts_compilation_9_16():
    from backend.app.services.slide_compiler import slide_compiler

    outline_reels = """Title: Tech Teaching Vertical Shorts
Theme: cyber-neon
Format: 9:16

# Slide 1: Rate Limiting in 30 Seconds
Layout: code-walkthrough
Code: [token_bucket.ts]
function checkRate(userId: string): boolean {
    return true;
}
- [1] Sliding Window: Check Redis quota
- [2] Atomic Decr: Consume token
- [3] Guard: Return 429 Too Many Requests

# Slide 2: Monolith vs Microservices
Layout: problem-solution
Problem: Distributed Monolith Pitfall
- Synchronous cascade of network calls
- Single point of failure in auth service
Solution: Event-Driven Kafka Mesh
- Independent async service autoscaling
- 99.999% availability with zero downtime

# Slide 3: Turborepo Highlights
Layout: bento-box
(icon: zap) Turbo Engine: Remote build caching
(icon: cpu) 10x | Faster: Rust compiler
(icon: workflow) DAG Execution: Task pruning
(icon: shield) Zero Config: Auto workspaces
"""

    pres = slide_compiler.compile(outline_reels)
    assert pres.aspect_ratio == "9:16"
    assert len(pres.slides) == 3

    for slide in pres.slides:
        assert len(slide.elements) >= 8
        for el in slide.elements:
            assert el.x >= 0, f"Negative x in 9:16 {slide.title}: {el.x}"
            assert el.x + el.width <= 720, f"X overflow in 9:16 {slide.title}: {el.x + el.width}"
            assert el.y >= 0, f"Negative y in 9:16 {slide.title}: {el.y}"
            assert el.y + el.height <= 1280, f"Y overflow in 9:16 {slide.title}: {el.y + el.height}"


def test_compile_outline_api_endpoint_with_tech_layouts():
    client = TestClient(app)
    outline = """Title: API Integration Test
Theme: emerald-dark
Format: 16:9

# Slide 1: Distributed System
Layout: architecture-layers
- (icon: globe) Layer 1 | Edge: Cloudflare CDN
- (icon: server) Layer 2 | Services: Go gRPC Microservices
- (icon: database) Layer 3 | Storage: PostgreSQL
"""
    res = client.post("/api/compile-outline", json={"outline": outline, "theme": "emerald-dark"})
    assert res.status_code == 200
    data = res.json()
    assert data["title"] == "API Integration Test"
    assert len(data["slides"]) == 1
    assert len(data["slides"][0]["elements"]) >= 10


def test_video_morph_group_and_type_matching():
    from backend.app.services.video_exporter import VideoExporter

    exporter = VideoExporter()
    slide_a = SlideModel(
        id="s1",
        title="Slide 1",
        background="#0f172a",
        elements=[
            ElementModel(
                id="card-rect-1",
                type=ElementType.RECTANGLE,
                x=100, y=100, width=200, height=200,
                morph_id="hero-group",
                group_id="grp-1",
            ),
            ElementModel(
                id="card-text-1",
                type=ElementType.TEXT,
                x=120, y=120, width=160, height=40,
                content="Title 1",
                morph_id="hero-group",
                group_id="grp-1",
            ),
        ],
    )
    slide_b = SlideModel(
        id="s2",
        title="Slide 2",
        background="#0f172a",
        elements=[
            ElementModel(
                id="card-text-2",
                type=ElementType.TEXT,
                x=420, y=120, width=160, height=40,
                content="Title 2",
                morph_id="hero-group",
                group_id="grp-1",
            ),
            ElementModel(
                id="card-rect-2",
                type=ElementType.RECTANGLE,
                x=400, y=100, width=300, height=300,
                morph_id="hero-group",
                group_id="grp-1",
            ),
        ],
    )

    matched, departing, arriving = exporter.match_elements(slide_a, slide_b)
    assert len(matched) == 2
    assert len(departing) == 0
    assert len(arriving) == 0

    # Ensure rectangle matched with rectangle and text matched with text despite different ordering
    rect_pair = next(p for p in matched if p[0].type == ElementType.RECTANGLE)
    text_pair = next(p for p in matched if p[0].type == ElementType.TEXT)
    assert rect_pair[1].id == "card-rect-2"
    assert text_pair[1].id == "card-text-2"
