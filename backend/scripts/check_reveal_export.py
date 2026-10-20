"""Inspect encoded reveal pixels and final-state PPTX without loading speech models."""
import json
from pathlib import Path
import subprocess
import sys

import imageio_ffmpeg
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))
from backend.app.models.presentation import PresentationModel
from backend.app.services.pptx_exporter import PPTXExporter

folder = ROOT / 'output/reveal-check'
ffmpeg = imageio_ffmpeg.get_ffmpeg_exe()
values = {}
for frame in [32, 33, 36, 39, 80]:
    pixels = subprocess.run([ffmpeg, '-v', 'error', '-i', str(folder / 'reveals.mp4'), '-vf',
        f'select=eq(n\\,{frame}),crop=20:20:240:240', '-frames:v', '1', '-f', 'rawvideo', '-pix_fmt', 'rgb24', '-'],
        capture_output=True, check=True).stdout
    values[frame] = sum(pixels[1::3]) / len(pixels[1::3])
assert values[32] < 3 and values[33] < 3, values
assert 78 < values[36] < 90 and 161 < values[39] < 176, values
assert abs(values[39] - values[80]) < 3, values
request = json.loads((folder / 'request.json').read_text())
deck = PresentationModel.model_validate(request['presentation'])
data = PPTXExporter().export(deck)
presentation = Presentation(data)
assert len(presentation.slides[0].shapes) >= 4  # Background plus all three elements.
assert str(presentation.slides[0].shapes[1].fill.fore_color.rgb) == '22D3EE'
report = dict(encoded_green_channel=values, expected_opacities=[0, 0, .4, .8, .8],
              final_state_pptx=True, model_inference=False, captions_enabled=False)
(folder / 'export-check.json').write_text(json.dumps(report, indent=2))
print(json.dumps(report, indent=2))
