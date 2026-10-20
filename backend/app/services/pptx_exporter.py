import io
import logging
from typing import Optional
import pptx
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls
from pptx.util import Inches, Pt
from .image_assets import rendered_image

from ..models.presentation import (
    PresentationModel,
    SlideModel,
    ElementModel,
    ElementType,
    ArrowHead,
    StrokeStyle,
)

logger = logging.getLogger(__name__)

CANVAS_WIDTH = 1280.0
CANVAS_HEIGHT = 720.0


def hex_to_rgb(hex_str: Optional[str]) -> Optional[RGBColor]:
    if not hex_str or hex_str == "transparent" or hex_str == "none":
        return None
    clean = hex_str.strip().lstrip("#")
    if len(clean) == 3:
        clean = "".join(c * 2 for c in clean)
    if len(clean) != 6:
        return None
    try:
        r = int(clean[0:2], 16)
        g = int(clean[2:4], 16)
        b = int(clean[4:6], 16)
        return RGBColor(r, g, b)
    except ValueError:
        return None


class PPTXExporter:
    def __init__(self):
        pass

    def export(self, presentation: PresentationModel) -> io.BytesIO:
        prs = pptx.Presentation()

        # Set slide dimensions
        if presentation.aspect_ratio == "9:16":
            slide_w_in = 7.5
            slide_h_in = 13.333
        elif presentation.aspect_ratio == "4:3":
            slide_w_in = 10.0
            slide_h_in = 7.5
        else:  # default 16:9 widescreen
            slide_w_in = 13.333
            slide_h_in = 7.5

        prs.slide_width = Inches(slide_w_in)
        prs.slide_height = Inches(slide_h_in)

        scale_x = slide_w_in / (720 if presentation.aspect_ratio == '9:16' else CANVAS_WIDTH)
        scale_y = slide_h_in / (1280 if presentation.aspect_ratio == '9:16' else CANVAS_HEIGHT)

        blank_layout = prs.slide_layouts[6]  # Blank slide

        for slide_model in presentation.slides:
            slide = prs.slides.add_slide(blank_layout)

            # Slide background
            bg_rgb = hex_to_rgb(slide_model.background or "#ffffff")
            if bg_rgb:
                bg_rect = slide.shapes.add_shape(
                    MSO_SHAPE.RECTANGLE,
                    Inches(0),
                    Inches(0),
                    Inches(slide_w_in),
                    Inches(slide_h_in),
                )
                bg_rect.fill.solid()
                bg_rect.fill.fore_color.rgb = bg_rgb
                bg_rect.line.fill.background()  # No border

            # Render elements
            sorted_elements = sorted(slide_model.elements, key=lambda e: e.z_index)
            for elem in sorted_elements:
                self._render_element(slide, elem, scale_x, scale_y)

            # Notes
            if slide_model.notes:
                notes_slide = slide.notes_slide
                notes_slide.notes_text_frame.text = slide_model.notes

            # Slide Transition (Morph, Fade, etc.)
            self._apply_transition(slide, slide_model)

        buffer = io.BytesIO()
        prs.save(buffer)
        buffer.seek(0)
        return buffer

    def _render_element(
        self, slide, elem: ElementModel, scale_x: float, scale_y: float
    ):
        try:
            if elem.type == ElementType.IMAGE:
                image = rendered_image(elem, elem.opacity)
                source = io.BytesIO()
                image.save(source, format='PNG')
                source.seek(0)
                picture = slide.shapes.add_picture(source, Inches(elem.x * scale_x), Inches(elem.y * scale_y),
                    width=Inches(max(1, elem.width) * scale_x), height=Inches(max(1, elem.height) * scale_y))
                picture.rotation = elem.rotation
            elif elem.type in (
                ElementType.RECTANGLE,
                ElementType.ROUNDED_RECTANGLE,
                ElementType.ELLIPSE,
            ):
                self._render_shape(slide, elem, scale_x, scale_y)
            elif elem.type == ElementType.TEXT:
                self._render_text(slide, elem, scale_x, scale_y)
            elif elem.type in (ElementType.LINE, ElementType.ARROW):
                self._render_line_or_arrow(slide, elem, scale_x, scale_y)
            elif elem.type == ElementType.ICON:
                self._render_icon(slide, elem, scale_x, scale_y)
            elif elem.type == ElementType.MOTION_COMPONENT:
                self._render_motion_component(slide, elem, scale_x, scale_y)
        except Exception as e:
            logger.error("Failed to render element %s (%s): %s", elem.id, elem.type, e)

    def _render_shape(
        self, slide, elem: ElementModel, scale_x: float, scale_y: float
    ):
        x = Inches(elem.x * scale_x)
        y = Inches(elem.y * scale_y)
        w = Inches(max(1.0, elem.width) * scale_x)
        h = Inches(max(1.0, elem.height) * scale_y)

        if elem.type == ElementType.ROUNDED_RECTANGLE:
            shape_type = MSO_SHAPE.ROUNDED_RECTANGLE
        elif elem.type == ElementType.ELLIPSE:
            shape_type = MSO_SHAPE.OVAL
        else:
            shape_type = MSO_SHAPE.RECTANGLE

        shape = slide.shapes.add_shape(shape_type, x, y, w, h)

        # Fill
        fill_rgb = hex_to_rgb(elem.fill)
        if fill_rgb:
            shape.fill.solid()
            shape.fill.fore_color.rgb = fill_rgb
        else:
            shape.fill.background()

        # Stroke / Border
        stroke_rgb = hex_to_rgb(elem.stroke)
        if stroke_rgb and elem.stroke_width > 0:
            shape.line.color.rgb = stroke_rgb
            shape.line.width = Pt(elem.stroke_width)
            if elem.stroke_style == StrokeStyle.DASHED:
                try:
                    ln = shape.line._get_or_add_ln()
                    ln.append(parse_xml(f'<a:prstDash val="dash" {nsdecls("a")}/>'))
                except Exception:
                    pass
        else:
            shape.line.fill.background()

        # Optional label content inside shape
        if elem.content:
            tf = shape.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = elem.content
            p.font.name = elem.font_family.split(",")[0].strip()
            p.font.size = Pt(max(8.0, elem.font_size * 0.75))
            p.font.bold = elem.font_weight == "bold"
            p.font.italic = elem.font_style == "italic"

            font_rgb = hex_to_rgb(elem.font_color)
            if font_rgb:
                p.font.color.rgb = font_rgb

            if elem.text_align == "center":
                p.alignment = PP_ALIGN.CENTER
            elif elem.text_align == "right":
                p.alignment = PP_ALIGN.RIGHT
            else:
                p.alignment = PP_ALIGN.LEFT

    def _render_text(
        self, slide, elem: ElementModel, scale_x: float, scale_y: float
    ):
        x = Inches(elem.x * scale_x)
        y = Inches(elem.y * scale_y)
        w = Inches(max(10.0, elem.width) * scale_x)
        h = Inches(max(10.0, elem.height) * scale_y)

        tx_box = slide.shapes.add_textbox(x, y, w, h)
        tf = tx_box.text_frame
        tf.word_wrap = True

        fill_rgb = hex_to_rgb(elem.fill)
        if fill_rgb:
            tx_box.fill.solid()
            tx_box.fill.fore_color.rgb = fill_rgb

        stroke_rgb = hex_to_rgb(elem.stroke)
        if stroke_rgb and elem.stroke_width > 0:
            tx_box.line.color.rgb = stroke_rgb
            tx_box.line.width = Pt(elem.stroke_width)

        lines = (elem.content or "").split("\n")
        for idx, line_text in enumerate(lines):
            p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
            p.text = line_text
            p.font.name = elem.font_family.split(",")[0].strip()
            p.font.size = Pt(max(8.0, elem.font_size * 0.75))
            p.font.bold = elem.font_weight == "bold"
            p.font.italic = elem.font_style == "italic"

            font_rgb = hex_to_rgb(elem.font_color)
            if font_rgb:
                p.font.color.rgb = font_rgb

            if elem.text_align == "center":
                p.alignment = PP_ALIGN.CENTER
            elif elem.text_align == "right":
                p.alignment = PP_ALIGN.RIGHT
            else:
                p.alignment = PP_ALIGN.LEFT

    def _render_line_or_arrow(
        self, slide, elem: ElementModel, scale_x: float, scale_y: float
    ):
        x1 = elem.x1 if elem.x1 is not None else elem.x
        y1 = elem.y1 if elem.y1 is not None else elem.y
        x2 = elem.x2 if elem.x2 is not None else (elem.x + elem.width)
        y2 = elem.y2 if elem.y2 is not None else (elem.y + elem.height)

        start_x = Inches(x1 * scale_x)
        start_y = Inches(y1 * scale_y)
        end_x = Inches(x2 * scale_x)
        end_y = Inches(y2 * scale_y)

        connector = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT, start_x, start_y, end_x, end_y
        )

        stroke_rgb = hex_to_rgb(elem.stroke or "#1e293b")
        if stroke_rgb:
            connector.line.color.rgb = stroke_rgb
        connector.line.width = Pt(max(1.0, elem.stroke_width))

        # Stroke style
        if elem.stroke_style == StrokeStyle.DASHED:
            try:
                ln = connector.line._get_or_add_ln()
                ln.append(parse_xml(f'<a:prstDash val="dash" {nsdecls("a")}/>'))
            except Exception:
                pass

        # Arrowheads
        if elem.type == ElementType.ARROW or elem.arrow_head != ArrowHead.NONE:
            try:
                ln = connector.line._get_or_add_ln()
                head = elem.arrow_head
                if head in (ArrowHead.END, ArrowHead.BOTH):
                    ln.append(
                        parse_xml(
                            f'<a:tailEnd type="triangle" w="med" len="med" {nsdecls("a")}/>'
                        )
                    )
                if head in (ArrowHead.START, ArrowHead.BOTH):
                    ln.append(
                        parse_xml(
                            f'<a:headEnd type="triangle" w="med" len="med" {nsdecls("a")}/>'
                        )
                    )
            except Exception as e:
                logger.warning("Could not set arrow head XML: %s", e)

    def _render_icon(
        self, slide, elem: ElementModel, scale_x: float, scale_y: float
    ):
        badge_style = getattr(elem, "icon_badge", "none") or "none"
        badge_color = getattr(elem, "badge_color", None) or elem.fill
        icon_name = getattr(elem, "icon_name", "box") or "box"
        label_text = elem.content.strip() if elem.content else ""

        x_in = Inches(elem.x * scale_x)
        y_in = Inches(elem.y * scale_y)
        w_in = Inches(max(1.0, elem.width) * scale_x)
        h_in = Inches(max(1.0, elem.height) * scale_y)

        badge_h_in = Inches(max(0.4, (elem.height - 24.0) * scale_y)) if label_text else h_in

        if badge_style == "circle":
            shape_type = MSO_SHAPE.OVAL
        elif badge_style == "rounded":
            shape_type = MSO_SHAPE.ROUNDED_RECTANGLE
        else:
            shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if badge_color else None

        if shape_type:
            shape = slide.shapes.add_shape(shape_type, x_in, y_in, w_in, badge_h_in)
            fill_rgb = hex_to_rgb(badge_color)
            if fill_rgb:
                shape.fill.solid()
                shape.fill.fore_color.rgb = fill_rgb
            else:
                shape.fill.background()

            stroke_rgb = hex_to_rgb(elem.stroke)
            if stroke_rgb and elem.stroke_width > 0:
                shape.line.color.rgb = stroke_rgb
                shape.line.width = Pt(elem.stroke_width)
            else:
                shape.line.fill.background()

            tf = shape.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = f"[{icon_name.upper()}]"
            p.font.name = "Arial"
            p.font.size = Pt(max(9.0, min(14.0, elem.width * scale_x * 12.0)))
            p.font.bold = True
            icon_rgb = hex_to_rgb(getattr(elem, "icon_color", None) or elem.stroke or "#ffffff")
            if icon_rgb:
                p.font.color.rgb = icon_rgb
            p.alignment = PP_ALIGN.CENTER

        if label_text:
            label_y_in = Inches((elem.y + elem.height - 22.0) * scale_y)
            label_box = slide.shapes.add_textbox(x_in, label_y_in, w_in, Inches(22.0 * scale_y))
            tf_lbl = label_box.text_frame
            tf_lbl.word_wrap = True
            p_lbl = tf_lbl.paragraphs[0]
            p_lbl.text = label_text
            p_lbl.font.size = Pt(max(8.0, elem.font_size * 0.75))
            lbl_rgb = hex_to_rgb(elem.font_color or "#334155")
            if lbl_rgb:
                p_lbl.font.color.rgb = lbl_rgb
            p_lbl.alignment = PP_ALIGN.CENTER

    def _render_motion_component(
        self, slide, elem: ElementModel, scale_x: float, scale_y: float
    ):
        motion = elem.motion_data
        card_color = (motion.card_color if motion and motion.card_color else "#0f172a")
        accent_color = (motion.accent_color if motion and motion.accent_color else "#38bdf8")
        text_color = (motion.text_color if motion and motion.text_color else "#ffffff")
        title = (motion.title if motion and motion.title else elem.content or "Motion Card")
        subtitle = (motion.subtitle if motion and motion.subtitle else "")
        tag = (motion.tag if motion and motion.tag else "")
        prefix = (motion.prefix if motion and motion.prefix else "")
        suffix = (motion.suffix if motion and motion.suffix else "")
        value_end = getattr(motion, "value_end", None) if motion else None

        x_in = Inches(elem.x * scale_x)
        y_in = Inches(elem.y * scale_y)
        w_in = Inches(max(1.0, elem.width) * scale_x)
        h_in = Inches(max(1.0, elem.height) * scale_y)

        # Base Card Shape (Rounded Rectangle)
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x_in, y_in, w_in, h_in)
        bg_rgb = hex_to_rgb(card_color)
        if bg_rgb:
            shape.fill.solid()
            shape.fill.fore_color.rgb = bg_rgb
        else:
            shape.fill.solid()
            shape.fill.fore_color.rgb = RGBColor(15, 23, 42)

        accent_rgb = hex_to_rgb(accent_color)
        if accent_rgb:
            shape.line.color.rgb = accent_rgb
            shape.line.width = Pt(1.5)
        else:
            shape.line.fill.background()

        # Text Frame inside the card
        tf = shape.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.15)
        tf.margin_right = Inches(0.15)
        tf.margin_top = Inches(0.15)
        tf.margin_bottom = Inches(0.15)

        # Top Tag / Badge
        p_tag = tf.paragraphs[0]
        has_tag = False
        if tag:
            p_tag.text = tag.upper()
            p_tag.font.name = "Arial"
            p_tag.font.size = Pt(8)
            p_tag.font.bold = True
            if accent_rgb:
                p_tag.font.color.rgb = accent_rgb
            has_tag = True

        # Metric value (if counting component)
        has_val = False
        if value_end is not None and (prefix or suffix or value_end > 0):
            val_str = f"{prefix}{value_end:g}{suffix}"
            p_val = tf.add_paragraph() if has_tag else p_tag
            p_val.text = val_str
            p_val.font.name = "Arial"
            p_val.font.size = Pt(max(16, min(28, int(elem.height * scale_y * 14))))
            p_val.font.bold = True
            if accent_rgb:
                p_val.font.color.rgb = accent_rgb
            has_val = True

        # Title
        p_title = tf.add_paragraph() if (has_tag or has_val) else p_tag
        p_title.text = title
        p_title.font.name = "Arial"
        p_title.font.size = Pt(max(10, min(14, int(elem.height * scale_y * 8))))
        p_title.font.bold = True
        txt_rgb = hex_to_rgb(text_color)
        if txt_rgb:
            p_title.font.color.rgb = txt_rgb

        # Subtitle
        if subtitle:
            p_sub = tf.add_paragraph()
            p_sub.text = subtitle
            p_sub.font.name = "Arial"
            p_sub.font.size = Pt(max(8, min(10, int(elem.height * scale_y * 6))))
            p_sub.font.color.rgb = RGBColor(148, 163, 184)

    def _apply_transition(self, slide, slide_model: SlideModel) -> None:
        """Inject native PowerPoint Morph or Fade transition XML into the slide."""
        trans_type = getattr(slide_model, "transition_type", "morph") or "morph"
        if trans_type == "none":
            return

        try:
            sld = slide._element
            duration_ms = int(getattr(slide_model, "transition_duration", 1.0) * 1000)

            if trans_type == "morph":
                xml_str = (
                    f'<p:transition xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
                    f'xmlns:p16="http://schemas.microsoft.com/office/powerpoint/2015/main" '
                    f'spd="med" advClick="1" dur="{duration_ms}">'
                    f'<p16:morph option="byObject"/>'
                    f'</p:transition>'
                )
            elif trans_type == "fade":
                xml_str = (
                    f'<p:transition xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
                    f'spd="med" advClick="1" dur="{duration_ms}">'
                    f'<p:fade/>'
                    f'</p:transition>'
                )
            else:
                return

            transition_elm = parse_xml(xml_str)
            # Find insertion point before timing or extLst if present
            inserted = False
            for i, child in enumerate(sld):
                if child.tag.endswith(("timing", "extLst")):
                    sld.insert(i, transition_elm)
                    inserted = True
                    break
            if not inserted:
                sld.append(transition_elm)
        except Exception as e:
            logger.warning("Could not apply slide transition XML: %s", e)


pptx_exporter = PPTXExporter()


